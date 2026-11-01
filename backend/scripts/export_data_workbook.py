"""Export every piece of data MIRAGE uses into one Excel workbook: data/MIRAGE_data_inventory.xlsx.

Models and thresholds, external datasets, demo samples and scenarios, seed data, the case database, test sets,
copilot training phrasings, lexicons, the live demo recording and every results file.

    python scripts/export_data_workbook.py
"""
import hashlib
import json
import sqlite3
import sys
from dataclasses import asdict, is_dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BACKEND = Path(__file__).resolve().parents[1]
ROOT = BACKEND.parent
sys.path.insert(0, str(BACKEND))

from app.assistant import catalog, eval_set  # noqa: E402
from app.fusion import response  # noqa: E402
from app.samples import SAMPLES_DIR, SCENARIOS, SPOKEN  # noqa: E402
from app.witnesses import pretrained, words  # noqa: E402
from scripts import field_data  # noqa: E402

OUT = ROOT / "data" / "MIRAGE_data_inventory.xlsx"
MAX_CELL = 2000  # embeddings and long JSON are cut; the full value stays in its source file
FONT = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1F3A5F")

wb = Workbook()
index: list[tuple[str, str, str]] = []  # sheet, what it holds, source


def cell(v):
    if v is None or isinstance(v, (int, float, bool)):
        return v
    if is_dataclass(v):
        v = asdict(v)
    if isinstance(v, (dict, list, tuple, set)):
        v = json.dumps(list(v) if isinstance(v, set) else v, ensure_ascii=False, default=str)
    s = ILLEGAL_CHARACTERS_RE.sub("", str(v))
    return s if len(s) <= MAX_CELL else s[:MAX_CELL] + " …(truncated)"


def sheet(name: str, what: str, source: str, header: list[str], rows) -> None:
    ws = wb.create_sheet(name[:31])
    ws.append(header)
    for r in rows:
        ws.append([cell(v) for v in r])
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
            c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 60)
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [min(len(str(c.value or "")), 70) for c in ws[get_column_letter(i)][1:200]])
        ws.column_dimensions[get_column_letter(i)].width = max(10, min(width + 2, 72))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    index.append((ws.title, what, source))


def flatten(d, prefix=""):
    """Nested results JSON -> (path, value) rows; lists of scalars stay on one row."""
    if isinstance(d, dict):
        for k, v in d.items():
            yield from flatten(v, f"{prefix}.{k}" if prefix else str(k))
    elif isinstance(d, list) and d and all(isinstance(x, (dict, list)) for x in d):
        for i, v in enumerate(d):
            yield from flatten(v, f"{prefix}[{i}]")
    else:
        yield prefix, d


def sha(p: Path) -> str:
    return hashlib.sha256(p.read_bytes()).hexdigest()


# ---- models -----------------------------------------------------------------------------------------------------------
sheet("Models", "Every pretrained detector: purpose, source repo, published accuracy, threshold and fusion weight",
      "app/witnesses/pretrained.py", ["key", "name", "source (Hugging Face / file)", "witness", "role", "published accuracy",
                                      "threshold", "fusion weight"],
      [(k, v[0], v[1], *(pretrained.CARDS.get(k, ("", "", ""))), pretrained.THRESHOLDS.get(k), pretrained.WEIGHTS.get(k))
       for k, v in pretrained.MODELS.items()])
sheet("Local model files", "Trained or fitted artefacts stored on this machine", "backend/models_store/",
      ["file", "size (KB)"], [(p.name, round(p.stat().st_size / 1024, 1)) for p in sorted((BACKEND / "models_store").iterdir())
                              if p.is_file()])

# ---- external datasets --------------------------------------------------------------------------------------------------
sheet("External datasets", "Public datasets sampled for the field test (fetched with scripts/field_data.py fetch)",
      "scripts/field_data.py", ["name", "dataset", "config", "split", "rows in split", "sampled", "label"],
      [("hemg", "Hemg/Deepfake-Audio-Dataset", "default", "train", 100, 50, "real + fake speech"),
       ("ljspeech", "DynamicSuperb/SpeechDetection_LJSpeech", "default", "test", 200, 30, "real + fake speech"),
       ("indic_hindi", "SPRINGLab/IndicTTS-Hindi", "default", "train", 11825, 30, "real Hindi speech"),
       ("google_tamil", "ylacombe/google-tamil", "female", "train", 2335, 20, "real Tamil speech"),
       ("google_tamil_m", "ylacombe/google-tamil", "male", "train", 2000, 20, "real Tamil speech"),
       ("librispeech", "openslr/librispeech_asr", "all", "test.clean", 2620, 40, "real English speech"),
       ("flickr", "nlphuji/flickr30k", "TEST", "test", 31014, 40, "real photographs")])
fieldtest = BACKEND / "var" / "fieldtest"
sheet("Field test files", "Every file in the field-test corpus (downloaded and generated)", "backend/var/fieldtest/",
      ["path", "folder", "size (KB)"],
      [(str(p.relative_to(fieldtest)).replace("\\", "/"), str(p.parent.relative_to(fieldtest)).replace("\\", "/"),
        round(p.stat().st_size / 1024, 1)) for p in sorted(fieldtest.rglob("*")) if p.is_file()] if fieldtest.exists() else [])
sheet("Field test generation", "Inputs used to generate synthetic field-test media", "scripts/field_data.py",
      ["kind", "key", "value"],
      [("edge-tts voice", k, v) for k, v in field_data.EDGE.items()]
      + [("spoken script", k, v) for k, v in field_data.SCRIPTS.items()]
      + [("SAPI voice", i, v) for i, v in enumerate(field_data.SAPI)]
      + [("SAPI text", i, v) for i, v in enumerate(field_data.SAPI_TEXT)]
      + [("SD-Turbo prompt", i, v) for i, v in enumerate(field_data.SD_PROMPTS)]
      + [("image laundering", i, v) for i, v in enumerate(field_data.IMG_LAUNDER)]
      + [("audio laundering", i, v) for i, v in enumerate(field_data.AUD_LAUNDER)])

# ---- demo data ------------------------------------------------------------------------------------------------------------
sheet("Demo samples", "Built-in sample evidence used by scenarios, the copilot and Live Shield", "backend/var/samples/",
      ["file", "size (KB)", "sha256"],
      [(p.name, round(p.stat().st_size / 1024, 1), sha(p)) for p in sorted(SAMPLES_DIR.iterdir()) if p.is_file()])
sheet("Scenarios", "Demo scenarios A–I with expected verdicts", "app/samples.py",
      list(asdict(SCENARIOS[0]).keys()), [list(asdict(s).values()) for s in SCENARIOS])
sheet("Spoken scripts", "Scripts the sample voice notes and video were generated from", "app/samples.py",
      ["key", "script", "voice"], [(k, t, v) for k, (t, v) in SPOKEN.items()])
demo = (ROOT / "frontend/src/pages/app/liveDemo.ts").read_text(encoding="utf-8")
calls = json.loads(demo.split("Record<string, DemoCall> = ", 1)[1].rstrip().rstrip(";"))
sheet("Live Shield demo", "Recorded per-window analysis replayed by the four Live Call Shield demo calls",
      "frontend/src/pages/app/liveDemo.ts",
      ["call", "window", "start (s)", "synthetic voice", "transcript", "tactics", "request", "manipulation", "amount (INR)"],
      [(f, i + 1, i * 3, w["spoof"], w["text"], ", ".join(w["tactics"]), w["action"], w["manipulation"], w["amount"])
       for f, c in calls.items() for i, w in enumerate(c["windows"])])

# ---- database ---------------------------------------------------------------------------------------------------------------
con = sqlite3.connect(BACKEND / "var" / "mirage.db")
tables = [t for (t,) in con.execute("select name from sqlite_master where type='table' order by name")]
for t in tables:
    cur = con.execute(f"select * from {t}")
    sheet(f"DB {t}", f"Every row of the '{t}' table (seed data, cases and activity)", "backend/var/mirage.db",
          [d[0] for d in cur.description], cur.fetchall())

# ---- test sets ------------------------------------------------------------------------------------------------------------
for name, f in (("Text cases (tuned)", "text_cases.json"), ("Text cases (held-out)", "text_cases_holdout.json")):
    rows = json.loads((ROOT / "testing" / f).read_text(encoding="utf-8"))
    sheet(name, f"Labelled scam / benign messages used to test the words witness ({len(rows)})", f"testing/{f}",
          list(rows[0].keys()), [list(r.values()) for r in rows])

# ---- copilot ---------------------------------------------------------------------------------------------------------------
sheet("Copilot intents", "The 52 intents the copilot router is trained on", "app/assistant/catalog.py",
      ["intent", "kind", "description", "training phrasings"],
      [(k, v.get("kind"), v.get("desc"), len(v.get("examples", []))) for k, v in catalog.INTENTS.items()])
sheet("Copilot phrasings", "Every training phrasing for the copilot router", "app/assistant/catalog.py",
      ["intent", "phrasing"], [(k, e) for k, v in catalog.INTENTS.items() for e in v.get("examples", [])])
sheet("Copilot eval set", "Held-out questions scoring the router (never used for training)", "app/assistant/eval_set.py",
      ["question", "expected intent"], eval_set.HELD_OUT)

# ---- lexicons and rules -------------------------------------------------------------------------------------------------------
sheet("Tactic lexicon", "Patterns the words witness uses to find pressure tactics (English regex and Indic words)",
      "app/witnesses/words.py", ["tactic", "language", "pattern"],
      [(k, "en (regex)", p) for k, v in words.TACTICS.items() for p in v]
      + [(k, "Indic", p) for k, v in words.INDIC_TACTICS.items() for p in v])
sheet("Entity patterns", "Regexes for UPI IDs, accounts, URLs, phones and other identifiers", "app/witnesses/words.py",
      ["entity", "pattern"], [(k, v.pattern) for k, v in words.ENTITY_RX.items()])
sheet("Helplines", "Reporting channels shown in the emergency response card", "app/fusion/response.py",
      ["id", "label", "link", "why"], [(k, v["label"], v["href"], v["why"]) for k, v in response.HELPLINES.items()])
sheet("Scam categories", "Categories a case is filed under", "app/fusion/response.py",
      ["id", "label"], list(response.CATEGORIES.items()))

# ---- results ---------------------------------------------------------------------------------------------------------------
for name, f in (("Results benchmark", BACKEND / "models_store/benchmark.json"),
                ("Results metrics", BACKEND / "models_store/metrics.json"),
                ("Results fusion", BACKEND / "models_store/fusion.json"),
                ("Results copilot eval", BACKEND / "models_store/assistant_eval.json"),
                ("Results field test", ROOT / "testing/results/field.json"),
                ("Results end-to-end", ROOT / "testing/results/e2e.json"),
                ("Results voice fusion", ROOT / "testing/results/voice_fusion.json"),
                ("Results voice candidates", ROOT / "testing/results/voice_candidates.json")):
    if f.exists():
        sheet(name, f"Every value in {f.name}, flattened (path = JSON key path)", str(f.relative_to(ROOT)).replace("\\", "/"),
              ["path", "value"], flatten(json.loads(f.read_text(encoding="utf-8"))))

# ---- index -----------------------------------------------------------------------------------------------------------------
ws = wb.active
ws.title = "Index"
ws.append(["MIRAGE data inventory: every dataset, sample, model, test set and result the project uses"])
ws.append(["Regenerate with: python backend/scripts/export_data_workbook.py. Cells over 2,000 characters are cut; "
           "the source file holds the full value."])
ws.append([])
ws.append(["sheet", "what it holds", "source in the codebase", "rows"])
for i, (name, what, src) in enumerate(index, start=5):
    ws.append([name, what, src, f"=MAX(COUNTA('{name}'!A:A)-1,0)"])
    ws.cell(i, 1).hyperlink = f"#'{name}'!A1"
ws.append(["Total", None, None, f"=SUM(D5:D{4 + len(index)})"])
ws["A1"].font = Font(name="Arial", size=13, bold=True)
ws["A2"].font = Font(name="Arial", size=9, italic=True, color="555555")
for c in ws[4]:
    c.font, c.fill = HEAD, FILL
for row in ws.iter_rows(min_row=5):
    for c in row:
        c.font = Font(name="Arial", size=10, color="0563C1", underline="single") if c.column == 1 and c.hyperlink else FONT
ws.cell(ws.max_row, 1).font = Font(name="Arial", size=10, bold=True)
ws.cell(ws.max_row, 4).font = Font(name="Arial", size=10, bold=True)
for col, w in zip("ABCD", (28, 90, 42, 10)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print(OUT, len(index), "sheets")
